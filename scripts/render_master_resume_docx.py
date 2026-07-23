from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from resume_format import date_range, humanize_group, training_suffix


ROOT = Path(__file__).resolve().parent.parent
GREEN = RGBColor(0x1F, 0x6F, 0x52)
MUTED = RGBColor(0x52, 0x60, 0x5A)


def text(value: object) -> str:
    return str(value or "")


def add_section_heading(document: Document, label: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(label)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = GREEN


def add_meta_line(document: Document, value: str) -> None:
    if not value:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(value)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def add_body(document: Document, value: str) -> None:
    if not value:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(value)
    run.font.size = Pt(10)


def add_bullet(document: Document, value: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(value)
    run.font.size = Pt(10)


def render(output: Path, data: dict) -> Path:
    career = data["career"]
    skills = data["skills"]
    technologies = data["technologies"]
    output.parent.mkdir(parents=True, exist_ok=True)
    contact = career["contact"]
    candidate = career["candidate"]

    document = Document()
    for style_name in ("Normal", "List Bullet"):
        document.styles[style_name].font.name = "Helvetica"

    name_paragraph = document.add_paragraph()
    name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_paragraph.add_run(text(candidate.get("name")))
    name_run.bold = True
    name_run.font.size = Pt(19)
    name_run.font.color.rgb = GREEN

    contact_paragraph = document.add_paragraph()
    contact_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_paragraph.paragraph_format.space_after = Pt(10)
    contact_run = contact_paragraph.add_run(" | ".join(text(contact.get(key)) for key in ("location", "email", "phone", "linkedin") if contact.get(key)))
    contact_run.font.size = Pt(9)
    contact_run.font.color.rgb = MUTED

    add_section_heading(document, "PROFESSIONAL SUMMARY")
    add_body(document, text(career["career_summary"]["positioning"]))

    add_section_heading(document, "WORK EXPERIENCE")
    for role in career.get("employment", []):
        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(1)
        run = heading.add_run(f"{text(role.get('official_title'))} | {text(role.get('organization'))}")
        run.bold = True
        run.font.size = Pt(10.5)
        add_meta_line(document, " | ".join(value for value in (text(role.get("location")), date_range(role.get("dates", {}))) if value))
        add_body(document, text(role.get("summary")))
        for item in role.get("responsibilities", []):
            add_bullet(document, text(item))

    add_section_heading(document, "PROJECTS")
    for project in career.get("projects", []):
        role = project.get("role")
        title = f"{text(project.get('name'))} | {text(role)}" if role else text(project.get("name"))
        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(1)
        run = heading.add_run(title)
        run.bold = True
        run.font.size = Pt(10.5)
        meta = " | ".join(value for value in (text(project.get("type")), text(project.get("organization")), text(project.get("status")), date_range(project.get("dates", {}))) if value)
        add_meta_line(document, meta)
        if project.get("url"):
            add_meta_line(document, text(project["url"]))
        add_body(document, text(project.get("summary")))
        for item in project.get("highlights", []):
            add_bullet(document, text(item))

    add_section_heading(document, "EDUCATION & CERTIFICATIONS")
    for item in career.get("education", []):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(f"{text(item.get('credential'))}, {text(item.get('field'))}")
        run.bold = True
        run.font.size = Pt(10)
        paragraph.add_run(f" - {text(item.get('institution'))} ({date_range(item.get('dates', {}))})").font.size = Pt(10)
    for item in career.get("training_and_certifications", []):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text(item.get("name")))
        run.bold = True
        run.font.size = Pt(10)
        paragraph.add_run(text(training_suffix(item))).font.size = Pt(10)

    add_section_heading(document, "SKILLS")
    for group, items in list(skills.items()) + list(technologies.items()):
        if not items:
            continue
        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(1)
        run = heading.add_run(humanize_group(group))
        run.bold = True
        run.font.size = Pt(9.5)
        body = document.add_paragraph()
        body.paragraph_format.space_after = Pt(4)
        body.add_run(", ".join(text(item.get("name")) for item in items)).font.size = Pt(9)

    document.save(str(output))
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", type=Path, required=True)
    args = parser.parse_args()
    path = render(ROOT / "output/pdf/Anna_Stupachenko_Master_Resume.docx", json.loads(args.data.read_text(encoding="utf-8")))
    print(path.relative_to(ROOT))
